"""Aplica correcciones de coherencia (cap. 1-3) sobre Avance_1_completo.docx."""
import copy, sys
import docx

PATH = "Avance_1_completo.docx"
d = docx.Document(PATH)
P = list(d.paragraphs)


def rep(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    raise SystemExit(f"no encontrado: {old[:50]}")


def set_text(p, new):
    p.runs[0].text = new
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find(start):
    m = [p for p in d.paragraphs if p.text.startswith(start)]
    assert len(m) == 1, (start, len(m))
    return m[0]


def cell_set(c, text):
    p = c.paragraphs[0]
    if p.runs:
        set_text(p, text)
    else:
        p.add_run(text)


def table_with(first_hdr, second_hdr=None):
    for t in d.tables:
        if t.rows[0].cells[0].text.startswith(first_hdr) and (second_hdr is None or t.rows[0].cells[1].text.startswith(second_hdr)):
            return t
    raise SystemExit("tabla no encontrada " + first_hdr)


# ---------------------------------------------- 3) espacios / alfabeto de 94
t = table_with("N.º", "Regla")
for r in t.rows:
    if r.cells[0].text == "R3":
        cell_set(r.cells[1], "Alfabeto: los 94 caracteres ASCII visibles (0x21 a 0x7E); el espacio no se toma en cuenta")
        cell_set(r.cells[2], "NIST: aceptar ASCII imprimible; el enunciado del proyecto indica que los espacios no se toman en cuenta")
    if r.cells[0].text == "R8":
        cell_set(r.cells[1], "Los espacios se descartan antes de validar (no forman parte del alfabeto)")
        cell_set(r.cells[2], "Enunciado del proyecto; NIST además permite recortar espacios en los extremos")
p = find("Este punto tiene consecuencia directa")
rep(p, "los 95 caracteres ASCII imprimibles, no solo letras y dígitos.",
    "los 94 caracteres ASCII visibles (0x21 a 0x7E), no solo letras y dígitos. El espacio (0x20) se descarta antes de leer, porque el enunciado del proyecto indica que los espacios no se toman en cuenta.")
rep(find("Σ  = { caracteres ASCII"), "{ caracteres ASCII imprimibles, 0x20 a 0x7E }   |Σ| = 95", "{ caracteres ASCII visibles, 0x21 a 0x7E }   |Σ| = 94")
t = table_with("Categoría")
for r in t.rows:
    if r.cells[0].text.strip() == "Espacio":
        cell_set(r.cells[2], "(se descarta al leer: el enunciado indica que los espacios no se toman en cuenta)")
        cell_set(r.cells[1], "0x20")
t = table_with("Dirección")
for r in t.rows:
    if r.cells[0].text.startswith("usuario ejemplo@"):
        cell_set(r.cells[2], "Espacio sin comillas en la parte local. El RFC lo invalida; el validador del capítulo 3 descarta los espacios (enunciado) y lo evalúa como usuarioejemplo@correo.com")
        cell_set(r.cells[1], "Inválida (RFC 5322)")

# ---------------------------------------------- 2) casos de prueba del cap. 1
t = table_with("Cadena de entrada")
for r in t.rows:
    s = r.cells[0].text
    if s.startswith("Mi perro"):
        cell_set(r.cells[2], "Los espacios se descartan (enunciado); queda «Miperrocome7Huesos!» de 19 caracteres con las cuatro clases")
    elif s.startswith("Password12345"):
        cell_set(r.cells[2], "Contiene la palabra prohibida «password» (P8) y la secuencia 2345 (P7); ambas son reglas léxicas, ver capítulo 3")
    elif s.startswith("Aaaa1111"):
        cell_set(r.cells[2], "Tres o más caracteres idénticos consecutivos (aaa, 1111, !!!!, bbb)")
    elif s.strip().startswith("Segura2026"):
        cell_set(r.cells[1], "Inválida")
        cell_set(r.cells[2], "Al descartar los espacios quedan 14 caracteres; incumple el mínimo de 15")

# ---------------------------------------------- 7) lookaheads
p71 = find("[\\x20-\\x7E]{15,64}$")
ref = find("La política de las reglas R1 a R7")
nota = copy.deepcopy(ref._p)
p71._p.addnext(nota)
from docx.text.paragraph import Paragraph
np_ = Paragraph(nota, p71._parent)
set_text(np_, "Nota: esta expresión con anticipaciones se incluye solo como referencia comparativa con los motores de expresiones "
              "regulares (PCRE, Python, JavaScript). Las anticipaciones no pertenecen al formalismo de ER que exige el punto 4 "
              "(unión, concatenación y cierre de Kleene), y su rango [\\x20-\\x7E] incluye el espacio, que el proyecto descarta. "
              "La ER formal, sobre el alfabeto de 94 símbolos, se desarrolla en el capítulo 3.")
for r in np_.runs:
    r.italic = True

# ---------------------------------------------- 5) párrafo pegado (cap. 2)
p174 = find("Es la que usan los navegadores")
tail = "Descomposición en tokens para el analizador léxico"
for r in p174.runs:
    if tail in r.text:
        r.text = r.text.replace(tail, "").rstrip()
h_ref = find("Autómata finito determinista para validar la dirección")
h_new = copy.deepcopy(h_ref._p)
p174._p.addnext(h_new)
set_text(Paragraph(h_new, p174._parent), tail)
p174.runs[-1].text = p174.runs[-1].text.rstrip()

# ---------------------------------------------- 4) AFD desactualizado (cap. 2)
p = find("El subconjunto de nivel 1 es un lenguaje regular")
set_text(p, "El subconjunto de nivel 1 es un lenguaje regular y admite un AFD. La siguiente formulación de siete estados es una "
            "versión simplificada: cubre la estructura (parte local, arroba y etiquetas del dominio) pero no verifica el dominio "
            "de nivel superior ni los límites de longitud.")
p = find("El estado de aceptación es q4")
set_text(p, "Limitación: q4 también acepta cadenas como usuario@ejemplo, sin punto ni TLD. El AFD completo, que exige al menos un "
            "punto después de la arroba y un TLD formado solo por letras (mínimo 2), se construye en el capítulo 3 (sección 3.6, "
            "Figura 6) con 12 estados por el método del árbol, 8 al minimizarlo.")

# ---------------------------------------------- 6) dibujo ASCII y líneas vacías
body = d.paragraphs
i0 = next(i for i, q in enumerate(body) if q.text.strip() == "PARTE LOCAL")
i1 = next(i for i, q in enumerate(body) if q.text.strip() == "Referencias" and i > i0)
for q in body[i0:i1]:
    q._p.getparent().remove(q._p)

# ---------------------------------------------- coherencia en el cap. 3
p = next(q for q in d.paragraphs if "Esto ajusta las reglas R3 y R8" in " ".join(r.text for r in q.runs))
rep(p, "Esto ajusta las reglas R3 y R8 del capítulo 1, que admitían espacio; aquí se sustituyen por el alfabeto de 94 símbolos ASCII visibles.",
    "Por eso las reglas R3 y R8 del capítulo 1 usan el alfabeto de 94 símbolos ASCII visibles, sin espacio.")
p = find("Nota: «␣»")
set_text(p, "Nota: «␣» marca un espacio en la entrada. La frase «Mi perro come 7 Huesos!» se evalúa como «Miperrocome7Huesos!» al "
            "descartar los espacios. «Password12345!!» se rechaza por P8 (palabra prohibida) y por P7 (secuencia 2345). Del mismo modo, "
            "«usuario ejemplo@correo.com» se evalúa como «usuarioejemplo@correo.com» y resulta válido, porque el enunciado indica "
            "que los espacios no se toman en cuenta; con las reglas estrictas de RFC 5322 sería inválido.")

d.save(PATH)
print("ok")
