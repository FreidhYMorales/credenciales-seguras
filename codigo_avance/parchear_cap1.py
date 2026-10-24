"""Puntos 3, 4 y 5 sobre el capitulo 1 de Avance_1_completo.docx."""
import copy
import docx
from docx.text.paragraph import Paragraph

d = docx.Document("Avance_1_completo.docx")


def set_text(p, s):
    p.runs[0].text = s
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find(start):
    m = [p for p in d.paragraphs if p.text.startswith(start)]
    assert len(m) == 1, (start, len(m))
    return m[0]


# 3. R10 alineada con P7 / P8
t = next(t for t in d.tables if t.rows[0].cells[0].text.startswith("N.º") and t.rows[0].cells[1].text == "Regla")
for r in t.rows:
    if r.cells[0].text == "R10":
        set_text(r.cells[1].paragraphs[0], "Sin secuencias ascendentes o descendentes de 4 símbolos consecutivos (abcd, dcba, 1234, 4321), sin distinguir mayúsculas")
        set_text(r.cells[2].paragraphs[0], "Reduce contraseñas predecibles. Palabras como qwerty se tratan como palabras prohibidas (P8, capítulo 3)")

# 5. reglas no verificables: distinguir blocklist real de lista finita
set_text(find("Comparación contra blocklist"),
         "Comparación contra la blocklist completa de contraseñas filtradas y palabras de diccionario (millones de entradas que cambian con el tiempo). "
         "Una lista finita y fija de palabras prohibidas (por ejemplo password, admin o qwerty) sí es un lenguaje regular y se implementa como regla P8 en el capítulo 3.")
set_text(find("Rechazo de términos del contexto"),
         "Rechazo de términos que dependen del contexto en tiempo de ejecución: nombre de usuario y derivados (varían con cada cuenta).")
p = find("Esta separación es exactamente")
set_text(p, p.text + " Por la misma razón, una lista finita y fija de palabras sí puede reconocerse con un autómata; lo que no puede es consultar una base externa y cambiante.")

# 4. alcance del autómata M y la gramática
anchor = find("Casos de prueba sugeridos")
ref = find("La gramática regular equivalente")
nota = copy.deepcopy(ref._p)
anchor._p.addprevious(nota)
np_ = Paragraph(nota, anchor._parent)
set_text(np_, "Alcance: la tupla M y la gramática anteriores cubren solo el alfabeto, la longitud y las cuatro clases de caracteres (R1 a R7) y están "
              "escritas sobre los 94 símbolos ASCII visibles. Los filtros de inseguridad se agregan como autómatas de rechazo, y el validador final es el "
              "autómata producto de todas las reglas. La correspondencia con el capítulo 3 es: R1 y R2 → P1; R3 → P0; R4 a R7 → P2 a P5; R9 → P6; "
              "R10 → P7; lista finita de palabras → P8. El diseño completo, con una expresión regular por regla, está en ese capítulo.")
for r in np_.runs:
    r.italic = True
d.save("Avance_1_completo.docx")
print("ok")
