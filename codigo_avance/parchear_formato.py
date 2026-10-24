"""Formato coherente: titulos, numeracion, tablas, bloques de codigo, referencias, pie de pagina."""
import copy
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

PATH = "Avance_1_completo.docx"
d = docx.Document(PATH)
body = d.element.body
BLACK = RGBColor(0, 0, 0)
CODE_INDENT = 228600


def P(el):
    return docx.text.paragraph.Paragraph(el, d._body)


def paragraphs():
    return [P(e) for e in body if e.tag == qn("w:p")]


def set_text(p, s):
    if not p.runs:
        p.add_run(s)
        return
    p.runs[0].text = s
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def style_run(r, size, bold=True, mono=False):
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = BLACK
    if mono:
        r.font.name = "Courier New"
        rpr = r._r.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), "Courier New")


def is_empty(el):
    if el.tag != qn("w:p"):
        return False
    x = el.xml
    return not P(el).text.strip() and "w:drawing" not in x and "w:br " not in x and "<w:br/>" not in x


def is_code(p):
    return p.style.name == "Normal" and p.paragraph_format.left_indent == CODE_INDENT and p.text.strip() != ""


# ---------------------------------------------------------------- 0. saltos de pagina sueltos
for p in paragraphs():
    if 'w:type="page"' in p._p.xml and not p.text.strip():
        p._p.getparent().remove(p._p)

# ---------------------------------------------------------------- 1. titulos de capitulo
CAP = re.compile(r"^[123] (Investigación|Diseño)")
for p in paragraphs():
    if CAP.match(p.text):
        p.style = d.styles["Heading 1"]
        set_text(p, p.text)
        style_run(p.runs[0], 14)
        pf = p.paragraph_format
        pf.page_break_before, pf.keep_with_next = True, True
        pf.space_before, pf.space_after = Pt(0), Pt(10)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

# ---------------------------------------------------------------- 2. referencias: unificar al final
ps = paragraphs()
ref_head = next(p for p in ps if p.text.strip() == "Referencias" and p.style.name == "Normal")
i = ps.index(ref_head)
entries = []
for p in ps[i + 1:]:
    if p.style.name.startswith("Heading"):
        break
    if p.text.strip():
        entries.append(p)
extra_head = next(p for p in ps if p.text.startswith("Referencias adicionales de este capítulo"))
for p in ps[ps.index(extra_head):ps.index(extra_head) + 3]:
    if p is extra_head or p.text.startswith(("Aho,", "Hopcroft,")):
        p._p.getparent().remove(p._p)
proto = copy.deepcopy(entries[0]._p)
for txt in (
    "Aho, A. V., Lam, M. S., Sethi, R. y Ullman, J. D. (2006). Compilers: Principles, Techniques, and Tools (2.ª ed.). Addison-Wesley.",
    "Hopcroft, J. E., Motwani, R. y Ullman, J. D. (2007). Introduction to Automata Theory, Languages, and Computation (3.ª ed.). Pearson/Addison-Wesley.",
):
    e = copy.deepcopy(proto)
    body.append(e)
    set_text(P(e), txt)
    entries.append(P(e))
entries.sort(key=lambda p: p.text.strip().lower())
sect = body.find(qn("w:sectPr"))
ref_head.style = d.styles["Heading 1"]
set_text(ref_head, "Referencias")
style_run(ref_head.runs[0], 14)
ref_head.paragraph_format.page_break_before = True
ref_head.paragraph_format.keep_with_next = True
ref_head.paragraph_format.space_after = Pt(10)
sect.addprevious(ref_head._p)
for e in entries:
    sect.addprevious(e._p)
    e.paragraph_format.space_after = Pt(6)

# ---------------------------------------------------------------- 3. subtitulos y numeracion
chap, k, old2new = None, 0, {}
for p in paragraphs():
    n = p.style.name
    if n == "Heading 1":
        m = re.match(r"^(\d) ", p.text)
        chap, k = (int(m.group(1)), 0) if m else (None, 0)
    elif n == "Heading 2":
        t = p.text
        m = re.match(r"^(\d+\.\d+)\s+(.*)$", t)
        if m:
            t = m.group(2)
        if chap:
            k += 1
            new = f"{chap}.{k}"
            if m:
                old2new[m.group(1)] = new
            t = f"{new} {t}"
        set_text(p, t)
        style_run(p.runs[0], 12)
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(12), Pt(4), True
    elif n == "Heading 3":
        set_text(p, p.text)
        style_run(p.runs[0], 11)
        pf = p.paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(8), Pt(3), True
for p in paragraphs():
    for r in p.runs:
        m = re.search(r"sección (3\.\d+)(?!\.\d)", r.text)
        if m and m.group(1) in old2new:
            r.text = r.text.replace(f"sección {m.group(1)},", f"sección {old2new[m.group(1)]},")
            print("ref cruzada:", m.group(1), "->", old2new[m.group(1)])

# ---------------------------------------------------------------- 4. texto: decimales y lineas de codigo largas
def set_code_lines(p, lines):
    set_text(p, lines[0])
    for ln in lines[1:]:
        p.runs[0].add_break()
        p.add_run(ln)
    for r in p.runs:
        style_run(r, 9, bold=False, mono=True)


for p in paragraphs():
    for r in p.runs:
        if "6,55" in r.text:
            r.text = r.text.replace("6,55", "6.55")
    t = p.text
    if t.startswith("P6  Σ*"):
        set_code_lines(p, ["P6  Σ* ( aaa | bbb | ... | ~~~ ) Σ*          94 alternativas xxx, x ∈ Σ"])
    elif t.startswith("P7  Σ*"):
        set_code_lines(p, ["P7  Σ* ( abcd | dcba | ... | 9876 ) Σ*        cada letra x se lee (x|X)"])
    elif t.startswith("P8  Σ*"):
        set_code_lines(p, ["P8  Σ* ( password | contrasena | ... ) Σ*     una alternativa por palabra"])
    elif t.startswith("COMP = Σ*LΣ*"):
        set_code_lines(p, ["COMP = Σ*LΣ*lΣ*NΣ*SΣ* | Σ*LΣ*lΣ*SΣ*NΣ* | ... | Σ*SΣ*NΣ*lΣ*LΣ*",
                           "       unión de las 4! = 24 permutaciones de L, l, N, S"])
    elif t.startswith("CORREO = at+"):
        set_code_lines(p, ["CORREO = at+ ( P at+ )* @ ( et P )+ A A+",
                           "         con  et = (A|N) ( (A|N|H)* (A|N) )?"])

# ---------------------------------------------------------------- 5. bloques de codigo uniformes
for p in paragraphs():
    if is_code(p):
        for r in p.runs:
            style_run(r, 9, bold=False, mono=True)
        pf = p.paragraph_format
        pf.space_after, pf.line_spacing, pf.left_indent = Pt(0), 1.0, Inches(0.25)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

# ---------------------------------------------------------------- 6. lineas vacias -> espaciado
kids = [e for e in body if e.tag != qn("w:sectPr")]
for idx, e in enumerate(kids):
    if not is_empty(e):
        continue
    prev = kids[idx - 1] if idx else None
    nxt = kids[idx + 1] if idx + 1 < len(kids) else None
    keep = (prev is not None and nxt is not None and
            ((prev.tag == qn("w:tbl") and nxt.tag == qn("w:tbl")) or
             (prev.tag == qn("w:p") and nxt.tag == qn("w:p") and is_code(P(prev)) and is_code(P(nxt)))))
    if not keep:
        e.getparent().remove(e)
kids = [e for e in body if e.tag != qn("w:sectPr")]
for idx, e in enumerate(kids):
    if e.tag != qn("w:p") or idx == 0:
        continue
    p, prev = P(e), kids[idx - 1]
    if (prev.tag == qn("w:tbl") or (prev.tag == qn("w:p") and is_code(P(prev)) and not is_code(p))) and not p.style.name.startswith("Heading"):
        p.paragraph_format.space_before = Pt(8)
    if idx + 1 < len(kids) and kids[idx + 1].tag == qn("w:tbl") and not p.style.name.startswith("Heading"):
        sa = p.paragraph_format.space_after
        if sa is None or sa < Pt(4):
            p.paragraph_format.space_after = Pt(4)

# ---------------------------------------------------------------- 7. tablas
def gridcols(t):
    return t._tbl.tblGrid.findall(qn("w:gridCol"))


TOTAL = 9360
for t in d.tables:
    has_pic = "w:drawing" in t._tbl.xml
    if has_pic:
        continue
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            trpr.append(OxmlElement("w:cantSplit"))
        for c in row.cells:
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(10)
    if len(t.rows) >= 5:
        hp = t.rows[0]._tr.get_or_add_trPr()
        if hp.find(qn("w:tblHeader")) is None:
            hp.append(OxmlElement("w:tblHeader"))
    cols = gridcols(t)
    ws = [int(c.get(qn("w:w"))) for c in cols]
    if len(set(ws)) == 1 and len(cols) > 1:  # ancho igual = sin diseñar: repartir segun contenido
        wt = []
        for j in range(len(cols)):
            m = max(len(row.cells[j].text) for row in t.rows)
            wt.append(min(max(m, 6), 55) ** 0.75)
        s = sum(wt)
        new = [int(TOTAL * w / s) for w in wt]
        new[-1] += TOTAL - sum(new)
        for c, w in zip(cols, new):
            c.set(qn("w:w"), str(w))
        for row in t.rows:
            for c, w in zip(row.cells, new):
                tcpr = c._tc.get_or_add_tcPr()
                tcw = tcpr.find(qn("w:tcW"))
                if tcw is None:
                    tcw = OxmlElement("w:tcW")
                    tcpr.insert(0, tcw)
                tcw.set(qn("w:w"), str(w))
                tcw.set(qn("w:type"), "dxa")

# ---------------------------------------------------------------- 8. pie de pagina con numero
sec = d.sections[0]
sec.different_first_page_header_footer = True
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in list(fp.runs):
    r._r.getparent().remove(r._r)
run = fp.add_run()
run.font.size = Pt(10)
for typ, txt in (("begin", None), (None, " PAGE "), ("end", None)):
    if typ:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), typ)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = txt
    run._r.append(el)

d.save(PATH)
print("ok")
