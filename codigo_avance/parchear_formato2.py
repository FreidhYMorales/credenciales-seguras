"""Relleno de celdas igual al de las tablas originales y sangria de 'Referencias'."""
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

d = docx.Document("docs/entregas/avance-1/Avance_1_completo.docx")
n = 0
for t in d.tables:
    if "w:drawing" in t._tbl.xml:
        continue
    for row in t.rows:
        for c in row.cells:
            tcpr = c._tc.get_or_add_tcPr()
            if tcpr.find(qn("w:tcMar")) is None:
                mar = OxmlElement("w:tcMar")
                for side, w in (("top", 60), ("left", 90), ("bottom", 60), ("right", 90)):
                    e = OxmlElement(f"w:{side}")
                    e.set(qn("w:w"), str(w))
                    e.set(qn("w:type"), "dxa")
                    mar.append(e)
                tcpr.append(mar)
                n += 1
for p in d.paragraphs:
    if p.style.name == "Heading 1":
        p.paragraph_format.left_indent = 0
        p.paragraph_format.first_line_indent = 0
d.save("docs/entregas/avance-1/Avance_1_completo.docx")
print("celdas ajustadas:", n)
