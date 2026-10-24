"""Puntos 6, 7 y 8: justificar P6-P8, limitacion de TLD xn-- y criterio de nivel de seguridad."""
import copy
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.text.paragraph import Paragraph

PATH = "docs/entregas/avance-1/Avance_1_completo.docx"
d = docx.Document(PATH)


def set_text(p, s):
    p.runs[0].text = s
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find(start):
    m = [p for p in d.paragraphs if p.text.startswith(start)]
    assert len(m) == 1, (start, len(m))
    return m[0]


def tbl(h0, h1):
    return next(t for t in d.tables if t.rows[0].cells[0].text == h0 and t.rows[0].cells[1].text.startswith(h1))


def cell(c, s):
    set_text(c.paragraphs[0], s)


# ---- 6. fundamentos de P6, P7, P8 (cap. 3) y R9/R10 (cap. 1)
t = tbl("ID", "Regla")
assert t.rows[0].cells[3].text == "Fundamento"
for r in t.rows:
    k = r.cells[0].text
    if k == "P6":
        cell(r.cells[3], "CIS §5.1.4: caracteres repetitivos o secuenciales (aaaaaa, 1234abcd)")
    elif k == "P7":
        cell(r.cells[3], "CIS §5.1.4: caracteres repetitivos o secuenciales (aaaaaa, 1234abcd)")
    elif k == "P8":
        cell(r.cells[1], "Contiene una palabra prohibida: password, contrasena, admin, qwerty, iloveyou, welcome, letmein (sin distinguir mayúsculas)")
        cell(r.cells[3], "CIS §5.1.4 (ejemplos: password, qwerty, iloveyou); NIST: blocklist obligatoria")
t = tbl("N.º", "Regla")
for r in t.rows:
    if r.cells[0].text == "R9":
        cell(r.cells[2], "Evita patrones triviales del tipo aaa o 111. CIS §5.1.4: caracteres repetitivos o secuenciales")
    if r.cells[0].text == "R10":
        cell(r.cells[2], "CIS §5.1.4: caracteres secuenciales (1234abcd). Palabras como qwerty se tratan como palabras prohibidas (P8, capítulo 3)")

p = find("La complejidad de la contraseña queda definida")
set_text(p, p.text + " Criterio de la lista de P8: password, qwerty e iloveyou aparecen como ejemplos de contraseñas prohibidas en la sección 5.1.4 del CIS "
                     "Password Policy Guide v21.12; contrasena es su equivalente en español; admin, welcome y letmein son ejemplos habituales de listas de "
                     "contraseñas comunes y se incluyen por criterio propio, no por cita del estándar. La lista es ampliable sin cambiar el método: cada palabra "
                     "nueva agrega una alternativa a la ER de P8.")

# tabla de tamanos: fila P8 con la lista ampliada (7 palabras)
t = tbl("ER", "Hojas")
for r in t.rows:
    if r.cells[0].text.startswith("P8"):
        for c, v in zip(r.cells[1:], ["59", "146", "213", "90", "44"]):
            cell(c, v)

# ---- 7. limitacion TLD xn--
for pgh in d.paragraphs:
    if pgh.text.startswith(" No se validan direcciones con comillas"):
        pass
p = next(q for q in d.paragraphs if "No se validan direcciones con comillas" in q.text)
r = p.runs[-1]
r.text = r.text + (" Tampoco se aceptan dominios de nivel superior internacionalizados codificados en Punycode (por ejemplo, xn--p1ai), "
                   "porque la regla C5 exige un TLD formado solo por letras.")

# ---- 8. criterio de nivel de seguridad (3.9), antes de las referencias adicionales
anchor = next(q for q in d.paragraphs if q.text.startswith("Referencias adicionales de este capítulo"))
ref_body = find("Este capítulo desarrolla")
ref_tbl = d.tables[0]
tbl_pr = copy.deepcopy(ref_tbl._tbl.tblPr)
hdr_shd = [copy.deepcopy(ch) for ch in ref_tbl.rows[0].cells[0]._tc.tcPr if ch.tag.endswith("}shd")]
body = d.element.body
before = set(body)


def para(text, bold=False, italic=False):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size = Pt(11)
    p.paragraph_format.space_after = ref_body.paragraph_format.space_after
    p.paragraph_format.line_spacing = ref_body.paragraph_format.line_spacing
    p.alignment = ref_body.alignment
    return p


h = d.add_paragraph(style="Heading 2")
hr = h.add_run("3.9 Criterio propuesto de nivel de seguridad (punto 7.a.iv)")
hr.bold, hr.font.size = True, Pt(11)
para("Este criterio es una propuesta de diseño propia para la etapa de la aplicación; no proviene de NIST ni de CIS. Se apoya en la idea de "
     "ambos estándares de que la longitud es el factor que más pesa, y se calcula sobre lo que ya reporta el validador.")
rows = [
    ["Baja", "Incumple alguna regla de aceptación (P0 a P5) o coincide con algún filtro de rechazo (P6 a P8)", "No es una contraseña aceptable"],
    ["Media", "Cumple todas las reglas y tiene de 15 a 19 caracteres", "Entre 98 y 124 bits (estimación bruta)"],
    ["Fuerte", "Cumple todas las reglas y tiene 20 o más caracteres", "131 bits o más (estimación bruta)"],
]
hdr = ["Nivel", "Condición", "Entropía bruta estimada"]
t = d.add_table(rows=1 + len(rows), cols=3)
t._tbl.remove(t._tbl.tblPr)
t._tbl.insert(0, copy.deepcopy(tbl_pr))
for j, x in enumerate(hdr):
    c = t.rows[0].cells[j]
    tp = c._tc.get_or_add_tcPr()
    for s in hdr_shd:
        tp.append(copy.deepcopy(s))
    rr = c.paragraphs[0].add_run(x)
    rr.bold, rr.font.size = True, Pt(9.5)
for i, row in enumerate(rows, 1):
    for j, x in enumerate(row):
        rr = t.rows[i].cells[j].paragraphs[0].add_run(x)
        rr.font.size = Pt(9.5)
para("La entropía bruta se estima como n · log2(94) ≈ 6,55 bits por símbolo, porque una contraseña que pasa todas las reglas contiene las cuatro clases "
     "y su alfabeto efectivo es de 94 símbolos. Es un límite superior, válido solo si los símbolos se eligieron al azar: una contraseña ideada por una persona "
     "tiene menos entropía real. Por eso el filtro de rechazo (P6 a P8) tiene prioridad y el nivel Baja se asigna antes de calcular nada.")
para("Ejemplos con el validador: Pl4t0n0#Guatemala2026 (21 caracteres) es Fuerte; Miperrocome7Huesos! (19) es Media; Xk9$mTq2#vLw8@Rz (16) es Media; "
     "Abc123! es Baja (incumple P1). Los umbrales de 15 y 20 caracteres son ajustables; 15 es el mínimo de NIST Rev. 4 para contraseñas de un solo factor, y "
     "20 es una marca práctica para frases largas. El generador del punto 7.a.iii podrá apuntar a longitudes de 20 o más para producir contraseñas Fuertes.")
new = [e for e in list(body) if e not in before and not e.tag.endswith("}sectPr")]
for e in new:
    anchor._p.addprevious(e)
d.save(PATH)
print("ok", len(new))
