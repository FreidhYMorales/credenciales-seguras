"""Cita precisa de PCI DSS v4.0.1 (req. 8.3.9, verificado en el PDF oficial) y su referencia."""
import copy
import docx
from docx.text.paragraph import Paragraph
d = docx.Document("docs/entregas/avance-1/Avance_1_completo.docx")
n = 0
for p in d.paragraphs:
    for r in p.runs:
        if r.text.startswith("Otro marco se alinea solo en parte"):
            r.text = ("Otro marco se alinea solo en parte con esta postura: el requisito 8.3.9 de PCI DSS v4.0.1 exige, cuando la contraseña es el único "
                      "factor de autenticación, cambiarla al menos cada 90 días, o bien analizar de forma dinámica la seguridad de la cuenta y decidir el "
                      "acceso en tiempo real; el requisito no aplica donde se usa autenticación multifactor. El requisito 8.3.6 pide al menos 12 caracteres "
                      "(8 si el sistema no admite 12) con letras y números.")
            n += 1
cis = next(p for p in d.paragraphs if p.text.startswith("Center for Internet Security. (2021)"))
new = copy.deepcopy(cis._p)
cis._p.addnext(new)
q = Paragraph(new, cis._parent)
q.runs[0].text = ("PCI Security Standards Council. (2024). Payment Card Industry Data Security Standard: Requirements and Testing Procedures, v4.0.1. "
                  "https://www.pcisecuritystandards.org")
for r in q.runs[1:]:
    r._r.getparent().remove(r._r)
assert n == 1
d.save("docs/entregas/avance-1/Avance_1_completo.docx")
print("ok")
